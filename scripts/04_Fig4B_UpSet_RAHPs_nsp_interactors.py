"""
Marano et al.

Figure 4B

Generate an UpSet-like plot showing the overlap of RAHPs with published
SARS-CoV-2 nsp interaction datasets grouped by viral protein function.

Input:
    data/RAHPs_nsp_interactions.xlsx

Output:
    figures/Fig4B_UpSet_RAHPs_nsp_interactors.pdf
"""

from pathlib import Path

import matplotlib.patches as mpatches
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]

DATA_DIR = ROOT / "data"
FIGURE_DIR = ROOT / "figures"

FIGURE_DIR.mkdir(exist_ok=True)

excel_file = DATA_DIR / "RAHPs_nsp_interactions.xlsx"

wb = openpyxl.load_workbook(excel_file)
ws = wb.active
headers = [cell.value for cell in ws[1]]

DMV_COLS = ["C: nsp3", "C: nsp4", "C: nsp6"]
RTC_COLS = ["C: nsp7", "C: nsp8", "C: nsp9", "C: nsp12"]
MRNA_COLS = ["C: nsp10", "C: nsp14", "C: nsp16", "C: nsp15", "C: nsp13"]

dmv_idx = [headers.index(n) for n in DMV_COLS]
rtc_idx = [headers.index(n) for n in RTC_COLS]
mrna_idx = [headers.index(n) for n in MRNA_COLS]
gene_idx = headers.index("T: GENE")

hits_dmv, hits_rtc, hits_mrna = set(), set(), set()
total = 0

for row in ws.iter_rows(min_row=2, values_only=True):
    gene = row[gene_idx]

    if not gene or "OC43" in gene.upper():
        continue

    total += 1

    if any(row[i] == "+" for i in dmv_idx):
        hits_dmv.add(gene)

    if any(row[i] == "+" for i in rtc_idx):
        hits_rtc.add(gene)

    if any(row[i] == "+" for i in mrna_idx):
        hits_mrna.add(gene)

only_dmv = hits_dmv - hits_rtc - hits_mrna
only_rtc = hits_rtc - hits_dmv - hits_mrna
only_mrna = hits_mrna - hits_dmv - hits_rtc
dmv_rtc = (hits_dmv & hits_rtc) - hits_mrna
dmv_mrna = (hits_dmv & hits_mrna) - hits_rtc
rtc_mrna = (hits_rtc & hits_mrna) - hits_dmv
all_three = hits_dmv & hits_rtc & hits_mrna

intersections = [
    ("DMV only", len(only_dmv), [True, False, False]),
    ("DMV + RTC", len(dmv_rtc), [True, True, False]),
    ("All three", len(all_three), [True, True, True]),
    ("RTC only", len(only_rtc), [False, True, False]),
    ("DMV + mRNA", len(dmv_mrna), [True, False, True]),
    ("mRNA only", len(only_mrna), [False, False, True]),
    ("RTC + mRNA", len(rtc_mrna), [False, True, True]),
]

set_sizes = [len(hits_dmv), len(hits_rtc), len(hits_mrna)]


def pct(n):
    return n / total * 100


COLOR_DMV = "#1D9E75"
COLOR_RTC = "#378ADD"
COLOR_MRNA = "#888780"
COLOR_DOT_INACTIVE = "#D8D8D8"


def bar_color(dmv, rtc, mrna):
    if dmv:
        return COLOR_DMV
    if rtc:
        return COLOR_RTC
    return COLOR_MRNA


plt.rcParams.update(
    {
        "font.family": "sans-serif",
        "font.size": 8,
        "axes.linewidth": 0.6,
        "xtick.major.width": 0.6,
        "ytick.major.width": 0.6,
    }
)

fig = plt.figure(figsize=(9, 5.5))
fig.patch.set_facecolor("white")

n_cols = len(intersections)
n_rows = 3

ax_bar = fig.add_axes([0.26, 0.44, 0.70, 0.48])
ax_dot = fig.add_axes([0.26, 0.16, 0.70, 0.26])
ax_set = fig.add_axes([0.03, 0.16, 0.20, 0.26])

bar_vals = [pct(x[1]) for x in intersections]
bar_cols = [bar_color(*x[2]) for x in intersections]
x_pos = np.arange(n_cols)

bars = ax_bar.bar(
    x_pos,
    bar_vals,
    color=bar_cols,
    width=0.55,
    zorder=3,
    linewidth=0,
    edgecolor="none",
)

for bar, value, (_, n, _) in zip(bars, bar_vals, intersections):
    x_center = bar.get_x() + bar.get_width() / 2
    y_center = bar.get_height()

    ax_bar.text(
        x_center,
        y_center + 0.3,
        f"{value:.1f}%",
        ha="center",
        va="bottom",
        fontsize=7.5,
        fontweight="bold",
        color="#111111",
    )

    ax_bar.text(
        x_center,
        y_center + 2.1,
        f"(n={n})",
        ha="center",
        va="bottom",
        fontsize=6.5,
        color="#555555",
    )

ax_bar.set_xlim(-0.5, n_cols - 0.5)
ax_bar.set_ylim(0, max(bar_vals) * 1.40)
ax_bar.set_xticks([])
ax_bar.set_ylabel("% of RAHPs", fontsize=8.5)
ax_bar.yaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:.0f}%"))
ax_bar.spines[["top", "right", "bottom"]].set_visible(False)
ax_bar.spines["left"].set_linewidth(0.6)
ax_bar.tick_params(axis="y", labelsize=7.5, length=3)
ax_bar.yaxis.grid(True, linestyle=":", linewidth=0.5, color="#CCCCCC", zorder=0)
ax_bar.set_axisbelow(True)

ax_bar.set_title(
    "% of RAHPs with SARS-CoV-2 nsp interactions\nby viral protein functional group",
    fontsize=9,
    fontweight="bold",
    pad=8,
    loc="left",
)

row_labels = [
    "mRNA processing\n(nsp10-13-14-15-16)",
    "RTC core\n(nsp7-8-9-12)",
    "DMV formation\n(nsp3-4-6)",
]

row_colors_dot = [COLOR_MRNA, COLOR_RTC, COLOR_DMV]

ax_dot.set_xlim(-0.5, n_cols - 0.5)
ax_dot.set_ylim(-0.5, n_rows - 0.5)
ax_dot.set_xticks([])
ax_dot.set_yticks(range(n_rows))
ax_dot.set_yticklabels(row_labels, fontsize=7.5)
ax_dot.tick_params(left=False, bottom=False)
ax_dot.spines[:].set_visible(False)

for i in range(n_rows):
    if i % 2 == 0:
        ax_dot.axhspan(i - 0.5, i + 0.5, color="#F7F7F7", zorder=0)

for col_i, (_, _, membership) in enumerate(intersections):
    mem_bottom_to_top = [membership[2], membership[1], membership[0]]
    active_rows = []

    for row_i, is_member in enumerate(mem_bottom_to_top):
        dot_color = row_colors_dot[row_i] if is_member else COLOR_DOT_INACTIVE

        ax_dot.scatter(
            col_i,
            row_i,
            s=100,
            color=dot_color,
            zorder=3,
            linewidths=0,
        )

        if is_member:
            active_rows.append(row_i)

    if len(active_rows) > 1:
        ax_dot.plot(
            [col_i, col_i],
            [min(active_rows), max(active_rows)],
            color=bar_color(*membership),
            linewidth=2.5,
            zorder=2,
            solid_capstyle="round",
        )

set_pcts = [pct(s) for s in set_sizes]
set_pcts_bottom_to_top = [set_pcts[2], set_pcts[1], set_pcts[0]]
set_colors_bottom_to_top = [COLOR_MRNA, COLOR_RTC, COLOR_DMV]
y_pos = np.arange(n_rows)

ax_set.barh(
    y_pos,
    set_pcts_bottom_to_top,
    color=set_colors_bottom_to_top,
    height=0.45,
    linewidth=0,
)

for y, value, n in zip(
    y_pos,
    set_pcts_bottom_to_top,
    [set_sizes[2], set_sizes[1], set_sizes[0]],
):
    ax_set.text(
        value + 0.8,
        y,
        f"{value:.0f}%  (n={n})",
        va="center",
        fontsize=6.5,
        color="#333333",
    )

max_pct = max(set_pcts_bottom_to_top)

ax_set.set_xlim(max_pct * 1.6, 0)
ax_set.set_ylim(-0.5, n_rows - 0.5)
ax_set.set_yticks([])
ax_set.set_xlabel("Set size (% RAHPs)", fontsize=7.5)
ax_set.xaxis.set_major_formatter(mticker.FuncFormatter(lambda v, _: f"{v:.0f}%"))
ax_set.spines[["top", "right", "left"]].set_visible(False)
ax_set.spines["bottom"].set_linewidth(0.6)
ax_set.tick_params(axis="x", labelsize=7, length=3)

for i in range(n_rows):
    if i % 2 == 0:
        ax_set.axhspan(i - 0.5, i + 0.5, color="#F7F7F7", zorder=0)

legend_elements = [
    mpatches.Patch(color=COLOR_DMV, label="DMV formation (nsp3-4-6)"),
    mpatches.Patch(color=COLOR_RTC, label="RTC core (nsp7-8-9-12)"),
    mpatches.Patch(color=COLOR_MRNA, label="mRNA processing (nsp10-13-14-15-16)"),
]

fig.legend(
    handles=legend_elements,
    loc="lower center",
    ncol=3,
    fontsize=7.5,
    frameon=False,
    bbox_to_anchor=(0.58, 0.00),
)

fig.text(
    0.97,
    0.01,
    f"n={total} RAHPs",
    ha="right",
    va="bottom",
    fontsize=7,
    color="#888888",
)

output_file = FIGURE_DIR / "Fig4B_UpSet_RAHPs_nsp_interactors.pdf"

fig.savefig(
    output_file,
    dpi=300,
    bbox_inches="tight",
    facecolor="white",
)

print(f"Saved: {output_file}")

plt.show()
